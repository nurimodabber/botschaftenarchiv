#!/usr/bin/env python3
"""
Bahá'í Messages Archive — Reformat Exported Messages
Re-formats all ~1,500 exported House messages in:
- Exportierte_Botschaften/01_Botschaften_des_Hauses
- Exportierte_Botschaften/02_Botschaften_an_oder_über_Institutionen

to match the exact official Universal House of Justice reference layout (Riḍván 2025):
- Page: A4, 1.0 inch (2.54 cm / 72 pt) margins on all sides
- Font: Palatino Linotype, 11pt, Pure Black (#000000)
- Header: DAS UNIVERSALE HAUS DER GERECHTIGKEIT / THE UNIVERSAL HOUSE OF JUSTICE (centered, uppercase, 11pt, regular, 48pt space after)
- Date: Riḍván YYYY (for Riḍván messages) or formatted German/English date (centered, 11pt, regular, 48pt space after)
- Recipient: Left-aligned, 11pt, regular, 24pt space after
- Salutation: Left-aligned, 11pt, regular, 14pt space after
- Numbered Paragraphs:
  - Paragraph number in 7.5pt at left margin, followed by tab
  - Tab stop at 0.5 inch (36pt)
  - Paragraph body text in 11pt, justified (Blocksatz)
  - Line spacing 1.15, space after 8pt
- Headings: Left-aligned, 11.5pt bold, 16pt space before, 6pt space after
- Blockquotes: Left indent 0.5 inch, right indent 0.25 inch, 10.5pt, justified
- Closing signature: Centered [gez.: Das Universale Haus der Gerechtigkeit] / [signed: ...] (48pt before)
- Footer: Centered page numbers on page 2+, first page footer empty
"""

import os
import sys
import re
import csv
import time
from pathlib import Path

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).resolve().parent.parent
# If running from scratch, resolve to workspace root
if not (BASE_DIR / "Exportierte_Botschaften").exists():
    BASE_DIR = Path("/Users/nurimilanmodabber/Library/CloudStorage/OneDrive-Personal/Baha'i/Botschaften Haus copy")

EXPORT_DIR = BASE_DIR / "Exportierte_Botschaften"
OVERVIEW_DIR = EXPORT_DIR / "00_Übersichten"
CSV_INDEX = OVERVIEW_DIR / "00_INHALTSVERZEICHNIS.csv"

COLOR_BLACK = RGBColor(0, 0, 0)
MONTHS_DE = ["", "Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]
MONTHS_EN = ["", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def set_palatino(run, size_pt, bold=False, italic=False):
    """Set font properties with full XML mapping to Palatino Linotype."""
    run.font.name = "Palatino Linotype"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = COLOR_BLACK
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Palatino Linotype')
    rFonts.set(qn('w:hAnsi'), 'Palatino Linotype')
    rFonts.set(qn('w:cs'), 'Palatino Linotype')
    rFonts.set(qn('w:eastAsia'), 'Palatino Linotype')
    rPr.append(rFonts)


def add_page_number(run):
    """Add dynamic Word PAGE field to run."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def setup_reference_page(doc):
    """Set exact margins, A4 geometry and clean footers matching reference layout."""
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # First page has no page number (matching reference PDF)
    section.different_first_page_header_footer = True

    # Subsequent pages have centered page number
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_pg = fp.add_run()
    set_palatino(run_pg, 10)
    add_page_number(run_pg)


def clean_typography(text):
    """Standardize Bahá'í diacritics and quotes."""
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    text = re.sub(r'Rid\s*\.?\s*v[aá]n', 'Riḍván', text, flags=re.IGNORECASE)
    text = re.sub(r'Bah[aá]\s*[\’\']\s*[ií]', 'Bahá’í', text, flags=re.IGNORECASE)
    text = re.sub(r'[\‘\']\s*Abdu[\’\']l\s*[-–]\s*Bah[aá]', '‘Abdu’l-Bahá', text, flags=re.IGNORECASE)
    text = re.sub(r'Bah[aá][\’\']u[\’\']ll[aá]h', 'Bahá’u’lláh', text, flags=re.IGNORECASE)
    text = re.sub(r'Shoghi\s+Effendi', 'Shoghi Effendi', text)
    return text.strip()


def process_body_lines(raw_lines):
    """Unwrap hard-wrapped lines and dehyphenate while preserving intentional breaks."""
    short_lines = [l for l in raw_lines if len(l) < 95]
    is_hard_wrapped = len(raw_lines) > 10 and (len(short_lines) / len(raw_lines)) > 0.6

    if not is_hard_wrapped:
        return [clean_typography(l) for l in raw_lines if l.strip()]

    paras = []
    curr = []

    for i, line in enumerate(raw_lines):
        line = line.strip()
        if not line:
            continue

        is_num_start = bool(re.match(r"^(\d+(\.\d+)?)\s+[A-ZÄÖÜa-zäöü“„\"'\(\]]", line))
        is_heading = len(line) < 70 and not line.endswith(".") and (line.isupper() or line.startswith("SECTION ") or line.startswith("TEIL ") or line.startswith("ABSCHNITT "))
        is_closing = bool(re.match(r"^(With\s+loving|Mit\s+(innigen|herzlichen))\b", line, re.I))

        if (is_num_start or is_heading or is_closing) and curr:
            paras.append(" ".join(curr))
            curr = []

        if curr and curr[-1].endswith("-"):
            prev = curr.pop()
            m_join = re.match(r"^[a-zäöü]", line)
            if m_join:
                first_word = line.split()[0]
                rest = line[len(first_word):].lstrip()
                curr.append(prev[:-1] + first_word)
                if rest:
                    curr.append(rest)
            else:
                curr.append(prev + line)
        else:
            curr.append(line)

        ends_punct = bool(re.search(r"[.!?:\u201d\u2019\"\']$", line))
        is_short = len(line) < 55
        next_line = raw_lines[i+1].strip() if i+1 < len(raw_lines) else ""
        next_is_upper = bool(re.match(r"^[A-ZÄÖÜ0-9“„\"]", next_line))

        if is_short and ends_punct and next_is_upper and not line.endswith(":"):
            if curr:
                paras.append(" ".join(curr))
                curr = []

    if curr:
        paras.append(" ".join(curr))

    return [clean_typography(p) for p in paras if p.strip()]


def reformat_single_document(meta, docx_path):
    """Reformat one DOCX file according to official reference layout."""
    try:
        old_doc = docx.Document(str(docx_path))
    except Exception as e:
        return False, f"Could not read docx: {e}"

    raw_lines = [p.text.strip() for p in old_doc.paragraphs if p.text.strip()]
    if not raw_lines:
        return False, "File is empty"

    lang = meta.get("language", "deutsch")
    is_de = (lang.lower() == "deutsch") or ("[DE]" in docx_path.name)
    source = meta.get("source", "UHG")
    is_itc = (source == "ITC") or ("ITC" in docx_path.name) or ("Lehrzentrum" in docx_path.name)
    doc_type = meta.get("type", "Botschaft")
    date_iso = meta.get("date", "")
    title = meta.get("title", "")

    # 1. Institution
    if is_itc:
        inst = "INTERNATIONALES LEHRZENTRUM" if is_de else "INTERNATIONAL TEACHING CENTRE"
        default_sig = "[gez.: Internationales Lehrzentrum]" if is_de else "[signed: International Teaching Centre]"
    else:
        inst = "DAS UNIVERSALE HAUS DER GERECHTIGKEIT" if is_de else "THE UNIVERSAL HOUSE OF JUSTICE"
        default_sig = "[gez.: Das Universale Haus der Gerechtigkeit]" if is_de else "[signed: The Universal House of Justice]"

    # 2. Date formatting
    is_ridvan = "ridvan" in doc_type.lower() or "riḍván" in doc_type.lower() or "ridvan" in title.lower() or "riḍván" in title.lower() or "ridvan" in docx_path.name.lower()
    is_nawruz = "naw-rúz" in doc_type.lower() or "naw-ruz" in doc_type.lower() or "naw-rúz" in title.lower() or "naw-ruz" in title.lower()

    m_date = re.search(r"^(\d{4})-(\d{2})-(\d{2})", docx_path.name)
    if not date_iso and m_date:
        date_iso = f"{m_date.group(1)}-{m_date.group(2)}-{m_date.group(3)}"

    if is_ridvan and date_iso and len(date_iso) >= 4:
        display_date = f"Riḍván {date_iso[:4]}"
    elif is_nawruz and date_iso and len(date_iso) >= 4:
        display_date = f"Naw-Rúz {date_iso[:4]}"
    elif date_iso and re.match(r"^\d{4}-\d{2}-\d{2}$", date_iso):
        y, m, d = date_iso.split('-')
        m_int, d_int = int(m), int(d)
        if 1 <= m_int <= 12:
            display_date = f"{d_int}. {MONTHS_DE[m_int]} {y}" if is_de else f"{d_int} {MONTHS_EN[m_int]} {y}"
        else:
            display_date = date_iso
    else:
        display_date = date_iso

    # 3. Parse letterhead elements from top
    idx = 0
    extracted_recipient = ""
    extracted_salutation = ""
    extracted_date_from_doc = ""

    while idx < min(10, len(raw_lines)):
        l = raw_lines[idx]

        # Skip solitary institution headers
        if len(l) < 55 and any(h in l.upper() for h in [
            "UNIVERSALE HAUS DER GERECHTIGKEIT", "UNIVERSAL HOUSE OF JUSTICE",
            "INTERNATIONALES LEHRZENTRUM", "INTERNATIONAL TEACHING CENTRE"
        ]):
            idx += 1
            continue

        # Check date pattern
        if re.match(r"^(\.?\d{1,2}\.?\s+[A-Za-zäöüÄÖÜ]+\s+\d{4}|Riḍván\s+\d{4}|\d{4}-\d{2}-\d{2})$", l, re.I):
            extracted_date_from_doc = l
            idx += 1
            continue

        # Skip exact title repetitions or isolated numbers
        if l == title or (len(title) > 20 and l.startswith(title[:35])) or re.match(r"^\d{1,4}$", l):
            idx += 1
            continue

        # Skip compilation header lines (e.g. '175 27 1 July 2013 To the participants...')
        if re.match(r"^\d+\s+\d+\s+\d{1,2}\s+[A-Za-z]+\s+\d{4}\s+To\s+", l):
            idx += 1
            continue

        # Check recipient
        if re.match(r"^(An\s+(die|alle|einen|das|ihre|seine|den|eine|Freunde)|To\s+(the|all|an|a\s+|the\s+Friends))\b", l, re.I):
            if not extracted_recipient:
                extracted_recipient = l
            idx += 1
            continue

        # Check salutation
        if re.match(r"^(Innig\s+geliebte|Dearly\s+loved|Dear\s+Bahá|Liebe\s+Bahá|Dear\s+Friends|Liebe\s+Freunde|Dear\s+Counsellors|Beloved\s+Friends|Liebe\s+Freundinnen)", l, re.I):
            if not extracted_salutation:
                extracted_salutation = l
            idx += 1
            continue

        break

    if not display_date and extracted_date_from_doc:
        display_date = extracted_date_from_doc

    # Recipient fallback from title
    if not extracted_recipient:
        m = re.search(r"[\(\[\–\-]\s*(An\s+[^)\]–-]+|To\s+[^)\]–-]+)[\)\]]?", title, re.IGNORECASE)
        if m:
            extracted_recipient = m.group(1).strip()
        elif re.match(r"^(An\s+(die|alle|einen|das|ihre|seine|den)|To\s+(the|all|an))\b", title, re.I):
            extracted_recipient = title.strip()

    # Clean bracketed recipient: '[To an individual]' -> 'To an individual'
    if extracted_recipient and extracted_recipient.startswith('[') and extracted_recipient.endswith(']'):
        extracted_recipient = extracted_recipient[1:-1].strip()

    # 4. Filter remaining lines
    body_raw = []
    closing_greeting = ""
    is_secretariat = False

    for l in raw_lines[idx:]:
        l_str = l.strip()
        if not l_str:
            continue

        # Check disclaimers and markers
        if any(bad in l_str for bad in ['Bahá’í Reference Library', 'terms of use found at', '-- Ende des Textes --']):
            continue

        # Check signature markers
        if re.match(r"^\s*\[?\s*(gez\.?:|signed:|gezeichnet:)\s*(Das\s+Universale|The\s+Universal|International)", l_str, re.I) or \
           l_str in ['[signed: The Universal House of Justice]', '[gez.: Das Universale Haus der Gerechtigkeit]']:
            continue

        if re.match(r"^(Department\s+of\s+the\s+Secretariat|Abteilung\s+des\s+Sekretariats)$", l_str, re.I):
            is_secretariat = True
            continue

        # Skip isolated small numbers (page numbers leaked from compilation)
        if re.match(r"^\d{1,4}$", l_str):
            continue

        # Skip compilation header lines
        if re.match(r"^\d+\s+\d+\s+\d{1,2}\s+[A-Za-z]+\s+\d{4}\s+To\s+", l_str):
            continue

        # Check closing greeting like 'With loving Bahá’í greetings,'
        if re.match(r"^(With\s+loving\s+Bahá’í\s+greetings|Mit\s+(innigen|herzlichen)\s+Bahá’í-Grüßen|With\s+warmest\s+Bahá’í\s+greetings|With\s+warm\s+Bahá’í\s+greetings)[,\.]?$", l_str, re.I):
            closing_greeting = l_str
            continue

        # Skip repeated recipient line inside body
        if extracted_recipient and (l_str == extracted_recipient or l_str == f"[{extracted_recipient}]"):
            continue

        body_raw.append(l_str)

    # 5. Unwrap lines into coherent paragraphs
    body_paras = process_body_lines(body_raw)

    # Determine final closing signature
    if is_secretariat:
        final_sig = "[Department of the Secretariat]" if not is_de else "[Abteilung des Sekretariats]"
    else:
        final_sig = default_sig

    # 6. BUILD DOCUMENT
    doc = docx.Document()
    setup_reference_page(doc)

    # Properties
    doc.core_properties.title = title[:240]
    doc.core_properties.author = inst
    doc.core_properties.subject = f"Bahá'í Guidance ({display_date})"[:240]

    # A. Header (Centered, 11pt, 48pt space after)
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(48)
    set_palatino(p_inst.add_run(inst), 11)

    # B. Date (Centered, 11pt, 48pt space after)
    if display_date:
        p_d = doc.add_paragraph()
        p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(48)
        set_palatino(p_d.add_run(display_date), 11)

    # C. Recipient (Left-aligned, 11pt, 24pt space after)
    if extracted_recipient:
        p_r = doc.add_paragraph()
        p_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_r.paragraph_format.space_before = Pt(0)
        p_r.paragraph_format.space_after = Pt(24)
        set_palatino(p_r.add_run(extracted_recipient), 11)

    # D. Salutation (Left-aligned, 11pt, 14pt space after)
    if extracted_salutation:
        p_s = doc.add_paragraph()
        p_s.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(14)
        set_palatino(p_s.add_run(extracted_salutation), 11)

    # E. Body Paragraphs (Numbered with 0.5 inch tab stop, justified, 1.15 spacing, 8pt after)
    p_num = 1
    for b_idx, p_text in enumerate(body_paras):
        p_clean = p_text.strip()
        if not p_clean:
            continue

        # Check if Heading
        if len(p_clean) < 75 and not p_clean.endswith('.') and (
            p_clean.isupper() or p_clean.startswith("SECTION ") or p_clean.startswith("TEIL ") or
            p_clean.startswith("ABSCHNITT ") or p_clean.startswith("KAPITEL ") or re.match(r"^[I|V|X]+\.\s+", p_clean)
        ):
            p_h = doc.add_paragraph()
            p_h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_h.paragraph_format.space_before = Pt(16)
            p_h.paragraph_format.space_after = Pt(6)
            set_palatino(p_h.add_run(p_clean), 11.5, bold=True)
            continue

        # Check if Blockquote
        prev_p = body_paras[b_idx - 1] if b_idx > 0 else ""
        is_quote = (p_clean.startswith("„") or p_clean.startswith("“") or p_clean.startswith('"')) and \
                   (p_clean.endswith("“") or p_clean.endswith("”") or p_clean.endswith('"') or prev_p.endswith(":"))

        if is_quote:
            p_q = doc.add_paragraph()
            p_q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_q.paragraph_format.left_indent = Inches(0.5)
            p_q.paragraph_format.right_indent = Inches(0.25)
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(8)
            p_q.paragraph_format.line_spacing = 1.15
            set_palatino(p_q.add_run(p_clean), 10.5)
            continue

        # Clean any existing leading numbers
        clean_body = p_clean
        m_num = re.match(r"^(\[(\d+)\]|(\d+)\.|\b(\d+)\b|(\d+\.\d+))\s+(.*)", p_clean)
        if m_num:
            clean_body = m_num.group(6) if m_num.group(6) else m_num.group(0)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))

        r_num = p.add_run(f"{p_num}\t")
        set_palatino(r_num, 7.5)

        r_txt = p.add_run(clean_body)
        set_palatino(r_txt, 11)
        p_num += 1

    # F. Closing greeting (if present)
    if closing_greeting:
        p_cg = doc.add_paragraph()
        p_cg.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cg.paragraph_format.space_before = Pt(18)
        p_cg.paragraph_format.space_after = Pt(24)
        set_palatino(p_cg.add_run(closing_greeting), 11)

    # G. Closing Signature (Centered, 48pt space before)
    p_cl = doc.add_paragraph()
    p_cl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cl.paragraph_format.space_before = Pt(48)
    p_cl.paragraph_format.space_after = Pt(12)
    set_palatino(p_cl.add_run(final_sig), 11)

    doc.save(str(docx_path))
    return True, f"Formatted ({p_num - 1} paras)"


def main():
    print("=" * 75)
    print("  Bahá'í Archive — Official Layout Re-formatter (Riḍván 2025 Standard)")
    print("=" * 75)

    # 1. Load CSV Index metadata lookup
    metadata_lookup = {}
    if CSV_INDEX.exists():
        with open(CSV_INDEX, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f, delimiter=";")
            for r in reader:
                rel = r.get("Relativer_Pfad", "").replace("../", "")
                if rel:
                    metadata_lookup[rel] = r
                fn = r.get("Dateiname", "")
                if fn:
                    metadata_lookup[fn] = r
        print(f"Loaded {len(metadata_lookup)} metadata mappings from index.")

    # 2. Collect target files
    target_dirs = [
        EXPORT_DIR / "01_Botschaften_des_Hauses",
        EXPORT_DIR / "02_Botschaften_an_oder_über_Institutionen"
    ]

    target_files = []
    for td in target_dirs:
        if td.exists():
            for f in td.rglob("*.docx"):
                target_files.append(f)

    print(f"Found {len(target_files)} DOCX files to format across sections 01 and 02.\n")

    t0 = time.time()
    success_count = 0
    fail_count = 0

    for idx, docx_file in enumerate(target_files):
        rel_key = os.path.relpath(docx_file, EXPORT_DIR)
        meta_row = metadata_lookup.get(rel_key) or metadata_lookup.get(docx_file.name) or {}

        meta = {
            "title": meta_row.get("Titel", docx_file.stem),
            "date": meta_row.get("Datum", ""),
            "year": meta_row.get("Jahr", ""),
            "language": meta_row.get("Sprache", "deutsch" if "[DE]" in docx_file.name else "english"),
            "source": meta_row.get("Institution", "UHG"),
            "type": meta_row.get("Dokumenttyp", "Botschaft")
        }

        ok, msg = reformat_single_document(meta, docx_file)
        if ok:
            success_count += 1
        else:
            fail_count += 1
            print(f"  [ERROR] {docx_file.name}: {msg}")

        if (idx + 1) % 150 == 0 or (idx + 1) == len(target_files):
            elapsed = time.time() - t0
            print(f"  Processed {idx + 1}/{len(target_files)} documents ({success_count} ok, {fail_count} failed) in {elapsed:.1f}s...")

    t1 = time.time()
    print("\n" + "=" * 75)
    print(f"  Completed formatting in {t1 - t0:.1f} seconds.")
    print(f"  Total succeeded: {success_count} / {len(target_files)}")
    print(f"  Total failed:    {fail_count}")
    print("=" * 75)


if __name__ == "__main__":
    main()
