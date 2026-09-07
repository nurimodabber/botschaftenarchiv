import os
import sys
import json
import html
import zipfile
import traceback
import argparse
import re
import shutil
from concurrent.futures import ProcessPoolExecutor, as_completed

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    BaseDocTemplate,
    PageTemplate,
    Frame,
    Paragraph,
    Spacer,
    KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_PATH = os.path.join(BASE_DIR, 'data', 'index.json')
TEXTS_DIR = os.path.join(BASE_DIR, 'data', 'texts')
FORMATS_DIR = os.path.join(BASE_DIR, 'documents', 'formats')

PDF_DIR = os.path.join(FORMATS_DIR, 'pdf')
DOCX_DIR = os.path.join(FORMATS_DIR, 'docx')
EPUB_DIR = os.path.join(FORMATS_DIR, 'epub')

os.makedirs(PDF_DIR, exist_ok=True)
os.makedirs(DOCX_DIR, exist_ok=True)
os.makedirs(EPUB_DIR, exist_ok=True)

# Register official Palatino fonts
font_path = '/System/Library/Fonts/Palatino.ttc'
if not os.path.exists(font_path):
    font_path = '/Library/Fonts/Palatino.ttc'

pdfmetrics.registerFont(TTFont('Palatino-Roman', font_path, subfontIndex=0))
pdfmetrics.registerFont(TTFont('Palatino-Italic', font_path, subfontIndex=1))
pdfmetrics.registerFont(TTFont('Palatino-Bold', font_path, subfontIndex=2))
pdfmetrics.registerFont(TTFont('Palatino-BoldItalic', font_path, subfontIndex=3))
pdfmetrics.registerFontFamily(
    'Palatino',
    normal='Palatino-Roman',
    bold='Palatino-Bold',
    italic='Palatino-Italic',
    boldItalic='Palatino-BoldItalic'
)

class OfficialHausCanvas(canvas.Canvas):
    """
    Canvas for official House of Justice documents.
    Draws centered page numbers at the bottom of all pages (1, 2, 3...)
    matching the official publication standard.
    """
    def showPage(self):
        self.saveState()
        self.setFont('Palatino-Roman', 11.04)
        # Position: 45.88 pt from bottom edge, centered on page width 595.28 pt
        self.drawCentredString(595.28 / 2.0, 45.88, str(self._pageNumber))
        self.restoreState()
        super().showPage()

class OfficialHausParagraph(Paragraph):
    """
    Paragraph with a small paragraph number drawn at the left margin (x=0 relative to frame)
    aligned with the baseline of the first line of text.
    The first line of text is indented by 36 pt (firstLineIndent=36.0),
    while subsequent lines return to the left margin.
    """
    def __init__(self, text, style, p_num=None, bulletText=None, frags=None, **kwargs):
        super().__init__(text, style, bulletText=bulletText, frags=frags, **kwargs)
        self.p_num = str(p_num) if p_num else None

    def draw(self):
        super().draw()
        if self.p_num:
            self.canv.saveState()
            self.canv.setFont('Palatino-Roman', 6.48)
            # Baseline alignment with line 0
            baseline_y = self.height - self.style.fontSize
            self.canv.drawString(0, baseline_y, self.p_num)
            self.canv.restoreState()

    def split(self, availWidth, availHeight):
        parts = super().split(availWidth, availHeight)
        if len(parts) > 0 and self.p_num:
            parts[0].p_num = self.p_num
            for p in parts[1:]:
                p.p_num = None
        return parts

def clean_xml(text):
    if not text:
        return ""
    clean = "".join(ch for ch in text if ch in ('\n', '\r', '\t') or ord(ch) >= 32)
    return html.escape(clean)

def normalize_paragraphs(raw_text):
    text = raw_text.replace('\r\n', '\n').replace('\r', '\n')
    lines = text.split('\n')
    new_lines = []
    for line in lines:
        stripped = line.strip()
        # If line starts with a paragraph number like '2 ', '3 ', '10 '
        if re.match(r'^\d+\s+[A-ZÄÖÜa-zäöü]', stripped) and not stripped.startswith(('19', '20')):
            new_lines.append('')
            new_lines.append(stripped)
        else:
            new_lines.append(stripped)
    norm = '\n'.join(new_lines)
    paras = [p.strip() for p in re.split(r'\n\s*\n+', norm) if p.strip()]
    return paras

def generate_pdf(doc, paragraphs, output_path):
    # A4 Dimensions: 595.28 x 841.89 pt
    # Left margin: 70.944 pt (~25 mm)
    # Right margin: 70.944 pt (~25 mm) -> Width = 453.392 pt
    # Top margin: 74.75 pt (~26.4 mm)
    # Bottom margin: 65.0 pt (~23 mm, above footer at 45.88 pt) -> Height = 702.14 pt
    frame = Frame(
        70.944, 65.0, 453.392, 702.14,
        id='MainFrame',
        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0
    )
    pdf_doc = BaseDocTemplate(output_path, pagesize=A4)
    pdf_doc.addPageTemplates([PageTemplate(id='Main', frames=frame)])

    header_style = ParagraphStyle(
        'HHead',
        fontName='Palatino-Roman',
        fontSize=12.0,
        leading=14.76,
        alignment=TA_CENTER,
        spaceAfter=46.0
    )

    date_style = ParagraphStyle(
        'HDate',
        fontName='Palatino-Roman',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_CENTER,
        spaceAfter=26.0
    )

    recipient_style = ParagraphStyle(
        'HRecip',
        fontName='Palatino-Roman',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_LEFT,
        spaceAfter=24.0
    )

    salutation_style = ParagraphStyle(
        'HSalut',
        fontName='Palatino-Roman',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_LEFT,
        spaceAfter=8.0
    )

    body_numbered_style = ParagraphStyle(
        'HBodyNum',
        fontName='Palatino-Roman',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_JUSTIFY,
        leftIndent=0,
        firstLineIndent=36.0,
        spaceAfter=8.0
    )

    body_plain_style = ParagraphStyle(
        'HBodyPlain',
        fontName='Palatino-Roman',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_JUSTIFY,
        leftIndent=0,
        firstLineIndent=36.0,
        spaceAfter=8.0
    )

    subheading_style = ParagraphStyle(
        'HSubhead',
        fontName='Palatino-Bold',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_LEFT,
        leftIndent=0,
        spaceBefore=12.0,
        spaceAfter=6.0
    )

    asterisk_style = ParagraphStyle(
        'HAst',
        fontName='Palatino-Roman',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_CENTER,
        spaceBefore=6.0,
        spaceAfter=6.0
    )

    sign_style = ParagraphStyle(
        'HSign',
        fontName='Palatino-Roman',
        fontSize=11.04,
        leading=14.76,
        alignment=TA_LEFT,
        leftIndent=180.0,
        spaceBefore=14.0,
        spaceAfter=8.0
    )

    closing_note_style = ParagraphStyle(
        'HClosing',
        fontName='Palatino-Roman',
        fontSize=10.0,
        leading=13.5,
        alignment=TA_LEFT,
        leftIndent=180.0,
        spaceBefore=10.0,
        spaceAfter=4.0
    )

    story = []
    in_header = True

    first_p = paragraphs[0] if paragraphs else ""
    first_upper = first_p.upper()
    has_explicit_header = any(h in first_upper for h in ['UNIVERSALE HAUS', 'UNIVERSAL HOUSE OF JUSTICE', 'INTERNATIONAL TEACHING CENTRE', 'LEHRZENTRUM'])

    lang = doc.get('language') or 'deutsch'
    is_itc = 'itc' in (doc.get('source') or '').lower()

    if not has_explicit_header:
        if is_itc:
            head_txt = 'DAS INTERNATIONALE LEHRZENTRUM' if lang == 'deutsch' else 'THE INTERNATIONAL TEACHING CENTRE'
        else:
            head_txt = 'DAS UNIVERSALE HAUS DER GERECHTIGKEIT' if lang == 'deutsch' else 'THE UNIVERSAL HOUSE OF JUSTICE'
        story.append(Paragraph(head_txt, header_style))

        date_val = doc.get('date') or ''
        # If date is not in first paragraph, add synthetic date
        if date_val and not any(date_val in p for p in paragraphs[:2]):
            story.append(Paragraph(clean_xml(date_val), date_style))

        recip = doc.get('recipientLabel') or doc.get('recipient')
        if recip and not any(recip in p for p in paragraphs[:3]):
            story.append(Paragraph(clean_xml(recip), recipient_style))

    para_counter = 0

    for idx, p_raw in enumerate(paragraphs):
        p_clean = p_raw.strip()
        if not p_clean:
            continue
        p_escaped = clean_xml(p_clean)

        p_upper = p_clean.upper()
        if any(p_upper == h for h in ['DAS UNIVERSALE HAUS DER GERECHTIGKEIT', 'THE UNIVERSAL HOUSE OF JUSTICE', 'DAS INTERNATIONALE LEHRZENTRUM', 'THE INTERNATIONAL TEACHING CENTRE']):
            story.append(Paragraph(p_escaped, header_style))
        elif in_header and (any(p_clean.startswith(kw) for kw in ['Riḍván', 'Naw-Rúz', '1', '2', '3', '4', '5', '6', '7', '8', '9', 'Jan', 'Feb', 'Mär', 'Mar', 'Apr', 'Mai', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Okt', 'Oct', 'Nov', 'Dez', 'Dec']) and len(p_clean) < 45):
            story.append(Paragraph(p_escaped, date_style))
        elif in_header and (p_clean.startswith(('An ', 'To ', 'Für ', 'Gegenüber ', 'An die ', 'To the ', 'To all ', 'An alle ')) and len(p_clean) < 140):
            story.append(Paragraph(p_escaped.replace('\n', '<br/>'), recipient_style))
        elif in_header and any(p_clean.endswith(s) for s in ['Freunde,', 'Friends,', 'Mitglieder,', 'Members,', 'Räte,', 'Councils,', 'Gläubigen,', 'Believers,', 'friend,', 'Freund,']):
            story.append(Paragraph(p_escaped.replace('\n', '<br/>'), salutation_style))
            in_header = False
        elif p_clean in ['*', '* * *', '***']:
            story.append(Paragraph('*', asterisk_style))
        elif p_clean.startswith(('[gez.:', '[signed:', '[gez.', '[signed')) or any(sig in p_clean for sig in ['Das Universale Haus der Gerechtigkeit]', 'The Universal House of Justice]']):
            story.append(Paragraph(p_escaped.replace('\n', '<br/>'), sign_style))
            in_header = False
        elif any(sec in p_clean for sec in ['Department of the Secretariat', 'Die Sekretariatsabteilung', 'Mit herzlichen Grüßen', 'With loving Bahá’í greetings']):
            story.append(Paragraph(p_escaped.replace('\n', '<br/>'), closing_note_style))
            in_header = False
        elif len(p_clean) < 70 and not p_clean.endswith(('.', ',', ':', ';', '!', '?')) and not re.match(r'^\d', p_clean) and not in_header:
            # Subheading
            story.append(Paragraph(p_escaped, subheading_style))
        else:
            in_header = False
            m = re.match(r'^(\d+(?:\.\d+)?)(?:[\.\t]|\s{2,}|\s+)(.*)$', p_clean, re.DOTALL)
            is_valid_p_num = False
            p_num = None
            body_text = p_escaped

            if m:
                num_str = m.group(1)
                try:
                    val = float(num_str)
                    if val < 500:
                        is_valid_p_num = True
                        p_num = num_str
                        body_text = clean_xml(m.group(2).strip())
                        para_counter = int(val)
                except ValueError:
                    pass

            if not is_valid_p_num:
                para_counter += 1
                p_num = str(para_counter)
                body_text = p_escaped

            story.append(OfficialHausParagraph(body_text, body_numbered_style, p_num=p_num))

    pdf_doc.build(story, canvasmaker=OfficialHausCanvas)

def generate_docx(doc, paragraphs, output_path):
    d = docx.Document()
    for sec in d.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    h = d.add_heading(doc.get('title', ''), level=1)
    if h.runs:
        h.runs[0].font.name = 'Palatino'
        h.runs[0].font.size = Pt(18)
        h.runs[0].font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    meta_p = d.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(12)
    date_str = doc.get('date', '')
    recipient_str = doc.get('recipientLabel') or doc.get('recipient') or ''
    meta_parts = [p for p in [date_str, doc.get('source', ''), recipient_str] if p]
    run_meta = meta_p.add_run(' • '.join(meta_parts))
    run_meta.font.name = 'Palatino'
    run_meta.font.italic = True
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    for p in paragraphs:
        clean_p = "".join(ch for ch in p if ch in ('\n', '\r', '\t') or ord(ch) >= 32)
        p_elem = d.add_paragraph(clean_p)
        p_elem.paragraph_format.line_spacing = 1.15
        p_elem.paragraph_format.space_after = Pt(6)
        p_elem.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if doc.get('sourceUrl'):
        src_p = d.add_paragraph()
        src_p.paragraph_format.space_before = Pt(18)
        run_src = src_p.add_run("Offizielle Referenz: " + doc['sourceUrl'])
        run_src.font.size = Pt(8.5)
        run_src.font.italic = True
        run_src.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

    d.save(output_path)

def generate_epub(doc, paragraphs, output_path):
    doc_id = doc.get('id', 'doc')
    title = clean_xml(doc.get('title', ''))
    source = clean_xml(str(doc.get('source') or "Bahá'í-Weltzentrum"))
    date_val = doc.get('date', '2026-01-01')
    lang = 'en' if doc.get('language') == 'english' else 'de'

    container_xml = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>"""

    content_opf = f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" unique-identifier="BookId" version="3.0">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:title>{title}</dc:title>
    <dc:creator>{source}</dc:creator>
    <dc:identifier id="BookId">urn:uuid:{doc_id}</dc:identifier>
    <dc:language>{lang}</dc:language>
    <dc:date>{date_val}</dc:date>
    <meta property="dcterms:modified">2026-09-07T00:00:00Z</meta>
  </metadata>
  <manifest>
    <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
    <item id="content" href="content.xhtml" media-type="application/xhtml+xml"/>
    <item id="style" href="style.css" media-type="text/css"/>
    <item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>
  </manifest>
  <spine toc="ncx">
    <itemref idref="content"/>
  </spine>
</package>"""

    nav_xhtml = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" lang="{lang}">
<head><title>{title}</title></head>
<body>
  <nav epub:type="toc" id="toc">
    <h1>Inhaltsverzeichnis</h1>
    <ol><li><a href="content.xhtml">{title}</a></li></ol>
  </nav>
</body>
</html>"""

    toc_ncx = f"""<?xml version="1.0" encoding="UTF-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
  <head><meta name="dtb:uid" content="urn:uuid:{doc_id}"/></head>
  <docTitle><text>{title}</text></docTitle>
  <navMap>
    <navPoint id="np-1" playOrder="1">
      <navLabel><text>{title}</text></navLabel>
      <content src="content.xhtml"/>
    </navPoint>
  </navMap>
</ncx>"""

    css = """body {
    font-family: Georgia, "Times New Roman", serif;
    margin: 5% 8%;
    line-height: 1.6;
    color: #111827;
    text-align: justify;
}
.eyebrow {
    font-size: 0.85em;
    text-transform: uppercase;
    letter-spacing: 0.12em;
    color: #c89d5c;
    margin-bottom: 0.4em;
    font-weight: bold;
}
h1 {
    font-size: 1.8em;
    line-height: 1.25;
    color: #0f172a;
    margin-top: 0;
    margin-bottom: 0.3em;
}
.meta {
    font-size: 0.9em;
    color: #64748b;
    font-style: italic;
    margin-bottom: 1.5em;
    border-bottom: 1.5px solid #c89d5c;
    padding-bottom: 0.8em;
}
p {
    margin-top: 0;
    margin-bottom: 0.9em;
    text-indent: 1.2em;
}
p.first-p {
    text-indent: 0;
}
.footer-note {
    margin-top: 3em;
    border-top: 1px solid #e2e8f0;
    padding-top: 1em;
    font-size: 0.8em;
    color: #94a3b8;
    font-style: italic;
}"""

    p_elems = []
    for i, p in enumerate(paragraphs):
        cls = ' class="first-p"' if i == 0 else ''
        p_elems.append(f'<p{cls}>{clean_xml(p)}</p>')
    p_html = '\n'.join(p_elems)

    recipient_str = doc.get('recipientLabel') or doc.get('recipient') or ''
    meta_parts = [p for p in [doc.get('date', ''), doc.get('source', ''), recipient_str] if p]
    meta_text = clean_xml(' • '.join(meta_parts))

    source_url_p = ''
    if doc.get('sourceUrl'):
        source_url_p = f'<p class="footer-note">Offizielle Referenz: {clean_xml(doc["sourceUrl"])}</p>'

    content_xhtml = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" lang="{lang}">
<head>
  <title>{title}</title>
  <link rel="stylesheet" type="text/css" href="style.css"/>
</head>
<body>
  <div class="eyebrow">{source.upper()}</div>
  <h1>{title}</h1>
  <div class="meta">{meta_text}</div>
  <div class="content">
    {p_html}
  </div>
  {source_url_p}
</body>
</html>"""

    with zipfile.ZipFile(output_path, 'w') as zf:
        zf.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
        zf.writestr('META-INF/container.xml', container_xml, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr('OEBPS/content.opf', content_opf, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr('OEBPS/nav.xhtml', nav_xhtml, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr('OEBPS/toc.ncx', toc_ncx, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr('OEBPS/style.css', css, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr('OEBPS/content.xhtml', content_xhtml, compress_type=zipfile.ZIP_DEFLATED)

def process_single_doc(args_tuple):
    doc, force_pdf = args_tuple
    doc_id = doc['id']
    txt_path = os.path.join(TEXTS_DIR, f"{doc_id}.txt")
    if not os.path.isfile(txt_path):
        return doc_id, False, f"Text file not found: {txt_path}"

    try:
        with open(txt_path, 'r', encoding='utf-8', errors='ignore') as f:
            raw_text = f.read().strip()

        paragraphs = normalize_paragraphs(raw_text)
        if not paragraphs:
            paragraphs = [doc.get('title', 'Dokument ohne Textinhalt.')]

        format_files = dict(doc.get('formatFiles') or {})
        available_fmts = set(doc.get('availableFormats') or [])

        # Ensure TXT
        format_files['txt'] = f"data/texts/{doc_id}.txt"
        available_fmts.add('txt')

        # 1. PDF
        target_pdf = os.path.join(PDF_DIR, f"{doc_id}.pdf")
        pdf_rel = format_files.get('pdf')
        needs_pdf = force_pdf or not pdf_rel or not os.path.isfile(os.path.join(BASE_DIR, pdf_rel))

        if needs_pdf or not os.path.isfile(target_pdf) or os.path.getsize(target_pdf) == 0:
            generate_pdf(doc, paragraphs, target_pdf)

            # If the doc had an existing legacy path (e.g. in documents/Buecher/...), overwrite it too
            if pdf_rel and pdf_rel != f"documents/formats/pdf/{doc_id}.pdf":
                legacy_full = os.path.join(BASE_DIR, pdf_rel)
                if os.path.isfile(legacy_full):
                    shutil.copy2(target_pdf, legacy_full)

            format_files['pdf'] = f"documents/formats/pdf/{doc_id}.pdf"

        available_fmts.add('pdf')

        # 2. DOCX
        docx_rel = format_files.get('docx')
        needs_docx = not docx_rel or not os.path.isfile(os.path.join(BASE_DIR, docx_rel))
        if needs_docx:
            target_docx = os.path.join(DOCX_DIR, f"{doc_id}.docx")
            if not os.path.isfile(target_docx) or os.path.getsize(target_docx) == 0:
                generate_docx(doc, paragraphs, target_docx)
            format_files['docx'] = f"documents/formats/docx/{doc_id}.docx"
        available_fmts.add('docx')

        # 3. EPUB
        epub_rel = format_files.get('epub')
        needs_epub = not epub_rel or not os.path.isfile(os.path.join(BASE_DIR, epub_rel))
        if needs_epub:
            target_epub = os.path.join(EPUB_DIR, f"{doc_id}.epub")
            if not os.path.isfile(target_epub) or os.path.getsize(target_epub) == 0:
                generate_epub(doc, paragraphs, target_epub)
            format_files['epub'] = f"documents/formats/epub/{doc_id}.epub"
        available_fmts.add('epub')

        canonical_order = ['pdf', 'docx', 'epub', 'txt']
        ordered_fmts = [f for f in canonical_order if f in available_fmts]
        for f in available_fmts:
            if f not in ordered_fmts:
                ordered_fmts.append(f)

        # Update file size based on primary format
        pdf_full_path = os.path.join(BASE_DIR, format_files.get('pdf', ''))
        file_size = os.path.getsize(pdf_full_path) if os.path.isfile(pdf_full_path) else doc.get('fileSize', 0)

        return doc_id, True, (format_files, ordered_fmts, file_size)
    except Exception as e:
        return doc_id, False, f"{str(e)}\n{traceback.format_exc()}"

def main():
    parser = argparse.ArgumentParser(description="Generate official PDF/DOCX/EPUB formats.")
    parser.add_argument('--house-only', action='store_true', help="Only process messages of the House of Justice (tiers house and institutions)")
    parser.add_argument('--force', action='store_true', help="Force regeneration of all PDFs")
    parser.add_argument('--doc-id', type=str, help="Process a single document by ID")
    parser.add_argument('--workers', type=int, default=0, help="Number of worker processes")
    args = parser.parse_args()

    print("Starting format generation...")
    with open(DATA_PATH, 'r', encoding='utf-8') as f:
        docs = json.load(f)

    if args.doc_id:
        target_docs = [d for d in docs if d['id'] == args.doc_id]
        if not target_docs:
            print(f"Document ID {args.doc_id} not found.")
            sys.exit(1)
    elif args.house_only:
        target_docs = [d for d in docs if d.get('tier') in ('house', 'institutions')]
    else:
        target_docs = docs

    print(f"Loaded {len(docs)} documents. Target documents to process: {len(target_docs)}.")

    success_count = 0
    error_count = 0
    doc_updates = {}

    workers = args.workers or min(os.cpu_count() or 4, 8)
    print(f"Using {workers} worker processes (force_pdf={args.force}).")

    work_items = [(doc, args.force) for doc in target_docs]

    with ProcessPoolExecutor(max_workers=workers) as executor:
        futures = {executor.submit(process_single_doc, item): item[0]['id'] for item in work_items}
        done = 0
        total = len(futures)
        for future in as_completed(futures):
            doc_id, ok, res = future.result()
            done += 1
            if ok:
                format_files, ordered_fmts, file_size = res
                doc_updates[doc_id] = (format_files, ordered_fmts, file_size)
                success_count += 1
            else:
                print(f"Error on {doc_id}: {res}")
                error_count += 1

            if done % 100 == 0 or done == total:
                print(f"Progress: {done}/{total} ({done*100/total:.1f}%) | Success: {success_count}, Errors: {error_count}")

    print("Updating index.json with new format metadata...")
    for doc in docs:
        doc_id = doc['id']
        if doc_id in doc_updates:
            ff, af, sz = doc_updates[doc_id]
            doc['formatFiles'] = ff
            doc['availableFormats'] = af
            if sz > 0:
                doc['fileSize'] = sz

    with open(DATA_PATH, 'w', encoding='utf-8') as f:
        json.dump(docs, f, ensure_ascii=False, indent=2)

    print(f"Finished! Successfully updated all {success_count} documents. Errors: {error_count}")

if __name__ == '__main__':
    main()
